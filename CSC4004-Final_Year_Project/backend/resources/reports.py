"""
DMS Resource Register Report Generator
Supports PDF, Excel (.xlsx), and CSV (.csv) formats.
"""
import csv
import io
import re
from datetime import date

INSTITUTION_LINE1 = "THE UNIVERSITY OF ZAMBIA"
INSTITUTION_LINE2 = "Department Of Computing & Informatics"
INSTITUTION_LINE3 = "Departmental Management System"
REPORT_TITLE = "RESOURCE REGISTER"

COLUMNS = [
    ("Resource ID", 14),
    ("Name", 28),
    ("Category", 20),
    ("Status", 12),
    ("Condition", 14),
    ("Location", 22),
    ("Portable", 10),
    ("Date Added", 14),
]

CONDITION_DISPLAY = {
    "NEW": "New", "GOOD": "Good", "FAIR": "Fair",
    "POOR": "Poor", "DAMAGED": "Damaged", "UNDER_REPAIR": "Under Repair",
}
STATUS_DISPLAY = {
    "AVAILABLE": "Available",
    "BOOKED": "Booked / Reserved",
    "IN_USE": "In Use",
    "UNAVAILABLE": "Unavailable",
    "MAINTENANCE": "Maintenance",
    "UNDER_REPAIR": "Under Repair",
    "RETIRED": "Retired",
}


def _resource_row(r):
    return [
        r.resource_id or "-",
        r.name,
        r.category.name if r.category else "Unassigned",
        STATUS_DISPLAY.get(r.status, r.status),
        CONDITION_DISPLAY.get(r.condition, r.condition),
        r.location or (r.assigned_room.name if r.assigned_room else "Unassigned"),
        "Yes" if r.is_portable else "No",
        str(r.date_added) if r.date_added else "-",
    ]


def report_filename(fmt, suffix="Resource_Register"):
    today = date.today().strftime("%Y-%m-%d")
    ext = {"PDF": "pdf", "EXCEL": "xlsx", "CSV": "csv"}.get(fmt.upper(), "pdf")
    safe_suffix = re.sub(r"[^A-Za-z0-9_-]+", "_", str(suffix)).strip("_") or "Resource_Register"
    return f"DMS_{safe_suffix}_{today}.{ext}"


def generate_pdf(resources, generated_by=None, category_filter=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.enums import TA_CENTER

    buffer = io.BytesIO()
    page_w, page_h = landscape(A4)
    margin = 18 * mm

    NAVY = colors.HexColor("#1e3a8a")
    LIGHT_GRAY = colors.HexColor("#f1f5f9")
    WHITE = colors.white
    ORANGE = colors.HexColor("#ea580c")
    MID_GRAY = colors.HexColor("#64748b")

    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin + 8 * mm,
        title="DMS Resource Register",
    )
    styles = getSampleStyleSheet()

    def _style(name, **kw):
        return ParagraphStyle(name, parent=styles["Normal"], **kw)

    s_title = _style("T", fontSize=15, fontName="Helvetica-Bold", textColor=NAVY, alignment=TA_CENTER, spaceAfter=2)
    s_sub = _style("S", fontSize=10, fontName="Helvetica", textColor=MID_GRAY, alignment=TA_CENTER, spaceAfter=2)
    s_report = _style("R", fontSize=13, fontName="Helvetica-Bold", textColor=ORANGE, alignment=TA_CENTER, spaceAfter=4)
    s_cell = _style("Cell", fontSize=8, leading=10, textColor=colors.HexColor("#0f172a"))
    s_head = _style("Head", fontSize=8, leading=9, fontName="Helvetica-Bold", textColor=WHITE, alignment=TA_CENTER)

    story = []
    story.append(Paragraph(INSTITUTION_LINE1, s_title))
    story.append(Paragraph(INSTITUTION_LINE2, s_sub))
    story.append(Paragraph(INSTITUTION_LINE3, s_sub))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1.5, color=NAVY))
    story.append(Spacer(1, 4))
    story.append(Paragraph(REPORT_TITLE, s_report))
    story.append(HRFlowable(width="100%", thickness=0.5, color=NAVY))
    story.append(Spacer(1, 6))

    today_str = date.today().strftime("%d %B %Y")
    meta_rows = [
        ["Generated:", today_str, "Total Resources:", str(len(list(resources)))],
        ["Category Filter:", category_filter or "All Categories",
         "Prepared by:", str(generated_by) if generated_by else "System"],
    ]
    resource_list = list(resources)
    meta_rows[0][3] = str(len(resource_list))

    meta_table = Table(meta_rows, colWidths=[28 * mm, 60 * mm, 38 * mm, 60 * mm])
    meta_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), MID_GRAY),
        ("TEXTCOLOR", (2, 0), (2, -1), MID_GRAY),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("FONTNAME", (3, 0), (3, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 8))

    usable_w = page_w - 2 * margin
    col_ratios = [c[1] for c in COLUMNS]
    total_r = sum(col_ratios)
    col_widths = [(r / total_r) * usable_w for r in col_ratios]

    headers = [Paragraph(c[0], s_head) for c in COLUMNS]
    data = [headers] + [[Paragraph(str(value), s_cell) for value in _resource_row(r)] for r in resource_list]

    t = Table(data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 8.5),
        ("ALIGN", (0, 1), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_GRAY]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, NAVY),
    ]
    for i, r in enumerate(resource_list, start=1):
        if r.condition in ("POOR", "DAMAGED"):
            style_cmds.append(("TEXTCOLOR", (4, i), (4, i), colors.HexColor("#dc2626")))
            style_cmds.append(("FONTNAME", (4, i), (4, i), "Helvetica-Bold"))
    t.setStyle(TableStyle(style_cmds))
    story.append(t)
    story.append(Spacer(1, 8))

    condition_counts = {}
    for r in resource_list:
        label = CONDITION_DISPLAY.get(r.condition, r.condition)
        condition_counts[label] = condition_counts.get(label, 0) + 1
    summary_text = "  |  ".join(f"{k}: {v}" for k, v in sorted(condition_counts.items()))
    story.append(Paragraph(
        f"Condition Summary - {summary_text}",
        _style("Sum", fontSize=8.5, textColor=MID_GRAY),
    ))

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(MID_GRAY)
        canvas.drawRightString(page_w - margin, 10 * mm, f"Page {doc.page}")
        canvas.drawString(margin, 10 * mm, f"{INSTITUTION_LINE3} - {REPORT_TITLE}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buffer.seek(0)
    return buffer.read()


def generate_excel(resources, generated_by=None, category_filter=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    resource_list = list(resources)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resource Register"

    NAVY = "1E3A8A"
    ORANGE = "EA580C"
    LGRAY = "F1F5F9"
    WHITE = "FFFFFF"
    MGRAY = "64748B"
    RED = "DC2626"

    thin = Side(style="thin", color="CBD5E1")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _cell(row, col, value, bold=False, color="000000", bg=None, align="left", size=11):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(bold=bold, color=color, size=size, name="Calibri")
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        return c

    ws.merge_cells("A1:H1")
    _cell(1, 1, INSTITUTION_LINE1, bold=True, color=NAVY, align="center", size=14)
    ws.merge_cells("A2:H2")
    _cell(2, 1, INSTITUTION_LINE2, color=MGRAY, align="center", size=11)
    ws.merge_cells("A3:H3")
    _cell(3, 1, INSTITUTION_LINE3, color=MGRAY, align="center", size=11)
    ws.merge_cells("A4:H4")
    _cell(4, 1, REPORT_TITLE, bold=True, color=ORANGE, align="center", size=13)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[4].height = 20

    today_str = date.today().strftime("%d %B %Y")
    ws.merge_cells("A6:B6")
    _cell(6, 1, "Generated:", bold=True, color=MGRAY)
    ws.merge_cells("C6:D6")
    _cell(6, 3, today_str)
    ws.merge_cells("A7:B7")
    _cell(7, 1, "Total Resources:", bold=True, color=MGRAY)
    ws.merge_cells("C7:D7")
    _cell(7, 3, len(resource_list))
    ws.merge_cells("A8:B8")
    _cell(8, 1, "Category Filter:", bold=True, color=MGRAY)
    ws.merge_cells("C8:D8")
    _cell(8, 3, category_filter or "All Categories")
    if generated_by:
        ws.merge_cells("A9:B9")
        _cell(9, 1, "Prepared by:", bold=True, color=MGRAY)
        ws.merge_cells("C9:D9")
        _cell(9, 3, str(generated_by))

    hr = 11
    for ci, (col_name, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=hr, column=ci, value=col_name)
        c.font = Font(bold=True, color=WHITE, size=10, name="Calibri")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = brd
    ws.row_dimensions[hr].height = 18

    for i, r in enumerate(resource_list):
        rn = hr + 1 + i
        row_data = _resource_row(r)
        bg = WHITE if i % 2 == 0 else LGRAY
        for ci, value in enumerate(row_data, 1):
            c = ws.cell(row=rn, column=ci, value=value)
            c.font = Font(size=10, name="Calibri")
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(vertical="center", horizontal="left")
            c.border = brd
            if ci == 5 and r.condition in ("POOR", "DAMAGED"):
                c.font = Font(bold=True, color=RED, size=10, name="Calibri")
        ws.row_dimensions[rn].height = 16

    for ci, w in enumerate([16, 30, 22, 14, 16, 24, 12, 16], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(COLUMNS))}{hr + len(resource_list)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_csv(resources):
    resource_list = list(resources)
    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([c[0] for c in COLUMNS])
    for r in resource_list:
        writer.writerow(_resource_row(r))
    return buf.getvalue().encode("utf-8-sig")
